import re
from dataclasses import dataclass
from pathlib import Path

import yaml


@dataclass
class ParsedDocument:
    title: str
    text: str
    source_path: str
    # The document's own link, from front matter. Shown to the user as the citation target.
    url: str | None = None
    # Page count, set for PDFs only: the preflight needs it to spot a scanned one.
    pages: int | None = None
    # `text` carries `#` headings, so the chunker may build a heading path from them.
    markdown: bool = False


def _title_from_markdown(text: str, path: Path) -> str:
    for line in text.splitlines():
        line = line.strip()
        if line.startswith("# "):
            return line[2:].strip()
    return path.stem.replace("-", " ").replace("_", " ").title()


_FRONT_MATTER = re.compile(r"\A---[ \t]*\n(.*?)\n---[ \t]*(?:\n|\Z)", re.DOTALL)


def parse_markdown(path: Path) -> ParsedDocument:
    text = path.read_text(encoding="utf-8")
    meta: dict = {}
    if m := _FRONT_MATTER.match(text):
        # Stripped from `text`: left in, it would be embedded and keyword-indexed as content.
        meta = yaml.safe_load(m.group(1)) or {}
        if not isinstance(meta, dict):
            raise ValueError("front matter must be a YAML mapping")
        text = text[m.end():].lstrip("\n")
    title = str(meta["title"]) if meta.get("title") else _title_from_markdown(text, path)
    url = str(meta["url"]) if meta.get("url") else None
    return ParsedDocument(
        title=title, text=text, source_path=str(path), url=url, markdown=True
    )


def parse_txt(path: Path) -> ParsedDocument:
    title = path.stem.replace("-", " ").replace("_", " ").title()
    return ParsedDocument(title=title, text=path.read_text(encoding="utf-8"), source_path=str(path))


def parse_docx(path: Path) -> ParsedDocument:
    import docx

    lines = []
    for para in docx.Document(str(path)).paragraphs:
        # "Heading 2" becomes "##" so the chunker's heading path works on Word files too.
        style = para.style.name if para.style is not None else ""
        level = style.removeprefix("Heading ")
        prefix = "#" * int(level) + " " if style.startswith("Heading ") and level.isdigit() else ""
        if para.text.strip():
            lines.append(prefix + para.text.strip())
    text = "\n\n".join(lines)
    return ParsedDocument(
        title=_title_from_markdown(text, path), text=text, source_path=str(path), markdown=True
    )


def parse_pdf(path: Path) -> ParsedDocument:
    import pypdf

    reader = pypdf.PdfReader(str(path))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    text = "\n\n".join(pages)
    title = path.stem.replace("-", " ").replace("_", " ").title()
    return ParsedDocument(title=title, text=text, source_path=str(path), pages=len(pages))


def parse_html(path: Path) -> ParsedDocument:
    from bs4 import BeautifulSoup

    html = path.read_text(encoding="utf-8")
    soup = BeautifulSoup(html, "html.parser")

    # Extract title from <title> or first <h1>
    title_tag = soup.find("title")
    h1_tag = soup.find("h1")
    if title_tag:
        title = title_tag.get_text(strip=True)
    elif h1_tag:
        title = h1_tag.get_text(strip=True)
    else:
        title = path.stem.replace("-", " ").replace("_", " ").title()

    # Remove script and style elements
    for tag in soup(["script", "style", "head"]):
        tag.decompose()

    text = soup.get_text(separator="\n")
    # Collapse excessive blank lines
    text = re.sub(r"\n{3,}", "\n\n", text).strip()

    return ParsedDocument(title=title, text=text, source_path=str(path))


_PARSERS = {
    ".md": parse_markdown,
    ".markdown": parse_markdown,
    ".txt": parse_txt,
    ".docx": parse_docx,
    ".pdf": parse_pdf,
    ".html": parse_html,
    ".htm": parse_html,
}


def parse_document(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    parser = _PARSERS.get(suffix)
    if parser is None:
        raise ValueError(f"Unsupported file type {suffix!r} for {path}")
    return parser(path)


def iter_corpus(corpus_dir: Path) -> list[Path]:
    """Return every file in a corpus directory, sorted. Which ones parse is the preflight's
    call to report, not something to decide here by leaving files out."""
    return sorted(p for p in corpus_dir.rglob("*") if p.is_file())


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in _PARSERS
