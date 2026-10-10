"""Preflight accounts for every file in the folder before anything is embedded or written,
and the real ingest uses the same report. Front matter, the heading path and the two new
fingerprint inputs are unit-tested alongside, on `tmp_path` alone.
"""
import hashlib
import json
import os
from pathlib import Path

import pypdf
import pytest
from sqlalchemy import delete, select

import app.ingest as ingest
from app.config.registry import ConfigRegistry
from app.config.schema import CorpusConfig
from app.ingest import preflight
from app.retrieval.chunker import chunk_document
from app.retrieval.parsers import parse_markdown

REPO_ROOT = Path(__file__).parents[3]
BODY = "Cloud Run scales instances on request concurrency. " * 12


def _cfg(tmp_path: Path, **corpus):
    cfg = ConfigRegistry(config_dir=REPO_ROOT / "config").get("gcp-platform-support")
    return cfg.model_copy(
        update={"client_id": "test-preflight", "corpus": CorpusConfig(source="corpus", **corpus)}
    )


def _folder(tmp_path: Path) -> Path:
    folder = tmp_path / "corpus"
    folder.mkdir()
    (folder / "good.md").write_text(f"# Good\n\n{BODY}\n")
    (folder / "sheet.xlsx").write_bytes(b"PK not read")
    (folder / "broken.pdf").write_bytes(b"%PDF-1.4\ntruncated")
    return folder


def _blank_pdf(path: Path) -> None:
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with path.open("wb") as f:
        writer.write(f)


def test_every_file_lands_in_exactly_one_column(tmp_path):
    _folder(tmp_path)
    report = preflight(_cfg(tmp_path), tmp_path)
    assert report.ingest == ["good.md"]
    assert report.unsupported == ["sheet.xlsx"]
    assert [f["path"] for f in report.failed] == ["broken.pdf"]
    assert len(report.files) == 3
    assert report.chunks > 0 and report.tokens > 0


def test_a_scanned_pdf_is_named_not_ingested_empty(tmp_path):
    folder = _folder(tmp_path)
    _blank_pdf(folder / "scan.pdf")
    report = preflight(_cfg(tmp_path), tmp_path)
    scan = next(f for f in report.failed if f["path"] == "scan.pdf")
    assert "needs OCR" in scan["reason"]


def test_two_files_with_one_text_are_flagged_and_a_stub_too(tmp_path):
    folder = _folder(tmp_path)
    (folder / "copy.txt").write_text(f"# Good\n\n{BODY}\n")
    (folder / "stub.txt").write_text("TODO")
    report = preflight(_cfg(tmp_path), tmp_path)
    assert report.duplicates == [["copy.txt", "good.md"]]
    assert report.stubs == ["stub.txt"]


def test_the_template_fills_stem_and_path_but_front_matter_wins(tmp_path):
    folder = _folder(tmp_path)
    (folder / "sub").mkdir()
    (folder / "sub" / "iam.md").write_text("# IAM\n\nx")
    (folder / "own.md").write_text("---\nurl: https://own.example/x\n---\n# Own\n\nx")
    cfg = _cfg(tmp_path, source_url_template="https://docs.acme.com/{path}#{stem}")
    urls = {k: v.url for k, v in preflight(cfg, tmp_path).parsed.items()}
    assert urls["sub/iam.md"] == "https://docs.acme.com/sub/iam#iam"
    assert urls["own.md"] == "https://own.example/x"
    assert _cfg(tmp_path).corpus.source_url_template == ""


def test_front_matter_is_stripped_and_honoured(tmp_path):
    f = tmp_path / "a.md"
    f.write_text('---\ntitle: "Real: title"\nurl: https://x.example/a\n---\n# Heading\n\nBody.\n')
    parsed = parse_markdown(f)
    assert (parsed.title, parsed.url) == ("Real: title", "https://x.example/a")
    assert "title:" not in parsed.text and parsed.text.startswith("# Heading")


def test_txt_and_docx_parse(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Setup", level=2)
    d.add_paragraph("Install it.")
    d.save(tmp_path / "w.docx")
    (tmp_path / "notes.txt").write_text("hello")
    from app.retrieval.parsers import parse_document

    assert parse_document(tmp_path / "w.docx").text.startswith("## Setup")
    assert parse_document(tmp_path / "notes.txt").title == "Notes"


def test_heading_path_follows_heading_changes():
    text = "# Doc\n\nintro\n\n## A\n\na body\n\n### A1\n\na1 body\n\n## B\n\nb body"
    chunks = chunk_document(
        text, "Doc", "u", chunk_size=4, overlap=0, prefix_title=True, markdown=True
    )
    heads = {c.text.split("\n")[0] for c in chunks}
    assert heads == {"Doc", "Doc — A", "Doc — A › A1", "Doc — B"}
    # Off by default: the text is untouched.
    plain = chunk_document(text, "Doc", "u", chunk_size=4, overlap=0, markdown=True)
    assert not any(c.text.startswith("Doc") and "—" in c.text for c in plain)
    # A non-markdown source gets the title alone, even if a line starts with "#".
    flat = chunk_document("## not a heading", "T", "u", prefix_title=True)
    assert flat[0].text == "T\n\n## not a heading"


def test_fingerprint_defaults_are_unchanged_and_new_keys_move_it():
    base = CorpusConfig(source="x")
    # The hash of the spec as it was before these keys existed; a change rebuilds every index.
    old = {"embedding_model": "voyage-3", "chunk_size": 800, "overlap": 100}
    assert base.fingerprint() == hashlib.sha256(json.dumps(old, sort_keys=True).encode()).hexdigest()[:16]
    assert base.fingerprint() == CorpusConfig(source="x", chunking={"prefix_title": False}).fingerprint()
    assert base.fingerprint() == CorpusConfig(source="x", source_url_template="").fingerprint()
    assert CorpusConfig(source="x", chunking={"prefix_title": True}).fingerprint() != base.fingerprint()
    assert CorpusConfig(source="x", source_url_template="https://a/{stem}").fingerprint() != base.fingerprint()


@pytest.mark.asyncio
async def test_check_embeds_nothing_and_writes_nothing(tmp_path, monkeypatch):
    async def boom(*a, **k):
        raise AssertionError("preflight must not embed")

    monkeypatch.setattr(ingest, "embed", boom)
    _folder(tmp_path)
    preflight(_cfg(tmp_path), tmp_path)  # no db is even in reach


@pytest.mark.skipif(os.getenv("RUN_INTEGRATION") != "1", reason="set RUN_INTEGRATION=1")
@pytest.mark.asyncio
async def test_a_bad_file_does_not_stop_the_good_one_committing(tmp_path, monkeypatch):
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
    from sqlalchemy.pool import NullPool

    from app.database import DATABASE_URL
    from app.models import Client, Document

    # Its own engine: the shared one's pooled connections belong to another test's event loop.
    AsyncSessionLocal = async_sessionmaker(
        create_async_engine(DATABASE_URL, poolclass=NullPool), expire_on_commit=False
    )

    async def fake_embed(texts, model):
        return [[0.0] * 1024 for _ in texts]

    monkeypatch.setattr(ingest, "embed", fake_embed)
    _folder(tmp_path)
    cfg = _cfg(tmp_path, source_url_template="https://docs.acme.com/{stem}")
    try:
        async with AsyncSessionLocal() as db:
            stats = await ingest.ingest_client(db, cfg, tmp_path)
            docs = (await db.execute(select(Document).where(Document.client_id == cfg.client_id))).scalars().all()
        assert stats["added"] == 1
        assert [f["path"] for f in stats["failed"]] == ["broken.pdf"]
        assert stats["report"].unsupported == ["sheet.xlsx"]
        assert [(d.title, d.url) for d in docs] == [("Good", "https://docs.acme.com/good")]
    finally:
        async with AsyncSessionLocal() as db:
            await db.execute(delete(Document).where(Document.client_id == cfg.client_id))
            await db.execute(delete(Client).where(Client.id == cfg.client_id))
            await db.commit()
